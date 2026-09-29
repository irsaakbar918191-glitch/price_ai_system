import csv
import io
import re
import requests
from PyPDF2 import PdfReader
import openpyxl
from docx import Document

class FileParser:
    @staticmethod
    def extract_raw_text(filename: str, file_bytes: bytes) -> str:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if ext == "csv":
            return FileParser._csv_to_text(file_bytes)
        elif ext in ("xlsx", "xls"):
            return FileParser._excel_to_text(file_bytes)
        elif ext == "pdf":
            return FileParser._pdf_to_text(file_bytes)
        elif ext in ("docx", "doc"):
            return FileParser._docx_to_text(file_bytes)
        elif ext in ("jpg", "jpeg", "png", "webp"):
            return FileParser._image_to_text(file_bytes, filename)
        else:
            return file_bytes.decode("utf-8", errors="ignore")

    @staticmethod
    def _csv_to_text(file_bytes: bytes) -> str:
        content = file_bytes.decode("utf-8", errors="ignore")
        reader = csv.reader(io.StringIO(content))
        lines = [" | ".join([str(c).strip() for c in r if str(c).strip()]) for r in reader if any(r)]
        return "\n".join(lines)

    @staticmethod
    def _excel_to_text(file_bytes: bytes) -> str:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for r in wb.active.iter_rows(values_only=True):
            cells = [str(c).strip() for c in r if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        return "\n".join(lines)

    @staticmethod
    def _pdf_to_text(file_bytes: bytes) -> str:
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])

    @staticmethod
    def _docx_to_text(file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for r in t.rows:
                cells = [c.text.strip() for c in r.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    @staticmethod
    def _image_to_text(file_bytes: bytes, filename: str) -> str:
        import base64
        b64_data = base64.b64encode(file_bytes).decode("utf-8")
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "jpeg"
        mime = f"image/{'png' if ext == 'png' else 'jpeg'}"

        try:
            payload = {
                "base64Image": f"data:{mime};base64,{b64_data}",
                "language": "eng",
                "isTable": "true",
                "scale": "true",
                "OCREngine": "2"
            }
            res = requests.post(
                "https://api.ocr.space/parse/image",
                data=payload,
                headers={"apikey": "K88289758888957"},
                timeout=25
            )
            if res.status_code == 200:
                parsed = res.json().get("ParsedResults", [])
                if parsed:
                    return parsed[0].get("ParsedText", "").strip()
        except Exception:
            pass

        return ""